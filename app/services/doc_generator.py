"""AI-powered resume and cover letter generator with DOCX export."""

import os
import json
from datetime import datetime, timezone
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from flask import current_app

from app import db
from app.models import Job, UserProfile, GeneratedDocument, ResumeVariant
from app.services.ai_client import ask


def _find_best_variant(job: Job, profile: UserProfile) -> ResumeVariant | None:
    """Find the resume variant whose target_roles best match the job title/description."""
    variants = ResumeVariant.query.filter_by(user_id=profile.id).all()
    if not variants:
        return None

    job_text = f"{job.title} {job.description or ''}".lower()
    best = None
    best_score = 0

    for v in variants:
        keywords = [k.strip().lower() for k in v.target_roles.split(",") if k.strip()]
        score = sum(1 for kw in keywords if kw in job_text)
        if score > best_score:
            best_score = score
            best = v

    return best if best_score > 0 else None


def generate_resume(job: Job, profile: UserProfile) -> GeneratedDocument:
    """Generate a tailored resume for a specific job posting."""
    system = (
        "You are an expert resume writer. Create a targeted, ATS-friendly resume "
        "tailored to the specific job posting. Use the candidate's actual experience "
        "and skills — do not fabricate anything. Emphasize the most relevant "
        "qualifications. Use strong action verbs and quantify achievements where possible."
    )

    variant = _find_best_variant(job, profile)
    variant_section = ""
    base_resume = profile.resume_text if profile.resume_text else "No existing resume provided."

    if variant and variant.resume_text:
        variant_section = f"""
## Targeted Resume Variant ({variant.label}):
{variant.resume_text[:4000]}
"""
        base_resume = variant.resume_text

    prompt = f"""Create a tailored resume for this candidate targeting the job below.

## Candidate Information:
- Name: {profile.name}
- Email: {profile.email}
- Phone: {profile.phone}
- Location: {profile.location}
- Skills: {profile.skills}
- Years of experience: {profile.experience_years}
- Summary: {profile.summary}

## Current Resume:
{base_resume}
{variant_section}

## Target Job:
- Title: {job.title}
- Company: {job.company}
- Description:
{job.description[:4000] if job.description else 'No description available'}

Format the resume using these sections with markdown:
# [Name]
[Contact info line]

## Professional Summary
[2-3 sentence targeted summary]

## Skills
[Relevant skills organized by category]

## Professional Experience
[Most relevant experience first, with bullet points using action verbs]

## Education
[Education details]

Only include information that's truthful based on the candidate's existing resume and profile.
Do NOT invent experience, degrees, or skills the candidate doesn't have.
"""

    content = ask(prompt, system=system)

    doc_record = GeneratedDocument(
        job_id=job.id,
        doc_type="resume",
        content=content,
    )
    db.session.add(doc_record)
    db.session.commit()

    # Generate DOCX
    filename = _generate_docx(doc_record, profile, job, "resume")
    doc_record.filename = filename
    db.session.commit()

    return doc_record


def generate_cover_letter(job: Job, profile: UserProfile) -> GeneratedDocument:
    """Generate a tailored cover letter for a specific job posting."""
    system = (
        "You are an expert cover letter writer. Create a compelling, personalized "
        "cover letter that connects the candidate's experience to the job requirements. "
        "Be professional but authentic. Do not fabricate experience."
    )

    variant = _find_best_variant(job, profile)
    base_resume = profile.resume_text if profile.resume_text else "No existing resume provided."
    sample_cover_letter_section = ""

    if variant:
        if variant.resume_text:
            base_resume = variant.resume_text
        if variant.cover_letter_text:
            sample_cover_letter_section = f"""
## Reference Cover Letter (use as style/tone guide — adapt content to this specific role):
{variant.cover_letter_text[:3000]}
"""

    prompt = f"""Write a cover letter for this candidate targeting the job below.

## Candidate Information:
- Name: {profile.name}
- Email: {profile.email}
- Phone: {profile.phone}
- Location: {profile.location}
- Summary: {profile.summary}

## Current Resume:
{base_resume}
{sample_cover_letter_section}

## Target Job:
- Title: {job.title}
- Company: {job.company}
- Location: {job.location}
- Description:
{job.description[:4000] if job.description else 'No description available'}

Write a professional cover letter (3-4 paragraphs) that:
1. Opens with enthusiasm for the specific role and company
2. Connects the candidate's most relevant experience to key job requirements
3. Highlights 2-3 specific achievements that demonstrate value
4. Closes with a confident call to action

Do NOT fabricate experience. Only reference skills and experience from the candidate's profile.
"""

    content = ask(prompt, system=system)

    doc_record = GeneratedDocument(
        job_id=job.id,
        doc_type="cover_letter",
        content=content,
    )
    db.session.add(doc_record)
    db.session.commit()

    filename = _generate_docx(doc_record, profile, job, "cover_letter")
    doc_record.filename = filename
    db.session.commit()

    return doc_record


def _generate_docx(
    doc_record: GeneratedDocument,
    profile: UserProfile,
    job: Job,
    doc_type: str,
) -> str:
    """Convert the text content to a DOCX file and return the filename."""
    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

    style = doc.styles["Normal"]
    style.font.size = Pt(11)
    style.font.name = "Calibri"

    # Parse markdown-ish content into the document
    lines = doc_record.content.split("\n")
    for line in lines:
        stripped = line.strip()
        if not stripped:
            doc.add_paragraph("")
            continue

        if stripped.startswith("# "):
            p = doc.add_heading(stripped[2:], level=1)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith("- ") or stripped.startswith("* "):
            doc.add_paragraph(stripped[2:], style="List Bullet")
        else:
            doc.add_paragraph(stripped)

    # Save file
    safe_company = "".join(c for c in job.company if c.isalnum() or c in " -_")[:30]
    safe_title = "".join(c for c in job.title if c.isalnum() or c in " -_")[:30]
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    filename = f"{doc_type}_{safe_company}_{safe_title}_{timestamp}.docx".replace(
        " ", "_"
    )

    folder = current_app.config["GENERATED_DOCS_FOLDER"]
    filepath = os.path.join(folder, filename)
    doc.save(filepath)

    return filename
